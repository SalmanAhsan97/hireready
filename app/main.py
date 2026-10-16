import hashlib
import hmac
import io
import os
import re
from pathlib import Path

from dotenv import load_dotenv

load_dotenv(Path(__file__).resolve().parent.parent / ".env")

from docx import Document  # noqa: E402
from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from . import ai, cache, jobs, pdf
from .models import Analysis, ChatMessage, Job
from .parsing import JobFetchError, extract_resume, extract_resume_text, fetch_job, fetch_linkedin_profile, linkedin_sources

STATIC = Path(__file__).resolve().parent.parent / "static"
MAX_UPLOAD = 5 * 1024 * 1024

app = FastAPI(title="HireReady")

# Public deploys: require a shared access code and cap paid work (scrapes + Claude calls) per visitor per day.
ACCESS_CODE = os.getenv("HIREREADY_ACCESS_CODE", "")
DAILY_LIMIT = int(os.getenv("HIREREADY_DAILY_LIMIT", "15"))  # analyses + application kits
CHAT_LIMIT = int(os.getenv("HIREREADY_CHAT_LIMIT", "40"))  # resume-chat messages
JOBS_LIMIT = int(os.getenv("HIREREADY_JOBS_LIMIT", "10"))  # job-board searches
_SALT = os.getenv("HIREREADY_CLIENT_SALT", ACCESS_CODE or "hireready")


PREVIEW_LIMIT = int(os.getenv("HIREREADY_PREVIEW_LIMIT", "3"))  # free previews per visitor per day
CONTACT_EMAIL = os.getenv("HIREREADY_CONTACT_EMAIL", "rnwedesign@gmail.com")


# Owner sign-in: full version with no daily limits. Credentials live only in the environment.
ADMIN_USER = os.getenv("HIREREADY_ADMIN_USER", "")
ADMIN_PASSWORD = os.getenv("HIREREADY_ADMIN_PASSWORD", "")
ADMIN_COOKIE = "hr_admin"
ADMIN_MAX_AGE = 30 * 24 * 3600
_login_failures: dict[str, int] = {}


def _admin_token() -> str:
    # changing the password invalidates every existing admin cookie
    return hmac.new(f"{_SALT}:{ADMIN_PASSWORD}".encode(), f"admin:{ADMIN_USER}".encode(), hashlib.sha256).hexdigest()


def is_admin(request: Request) -> bool:
    return bool(ADMIN_USER and ADMIN_PASSWORD) and hmac.compare_digest(request.cookies.get(ADMIN_COOKIE, ""), _admin_token())


def has_full_access(request: Request) -> bool:
    return (not ACCESS_CODE or is_admin(request)
            or hmac.compare_digest(request.headers.get("x-access-code", ""), ACCESS_CODE))


def require_access(request: Request):
    if not has_full_access(request):
        raise HTTPException(401, {
            "code": "access_code",
            "message": f"This is part of the full version. Email {CONTACT_EMAIL} for access and pricing, or enter your access code.",
        })


@app.middleware("http")
async def plan_for_request(request: Request, call_next):
    # paid requests get the premium model for writing; scoring is identical on every plan
    ai.set_plan(has_full_access(request))
    return await call_next(request)


def preview_of(a: Analysis) -> tuple[Analysis, dict]:
    """Free preview: the score and a taste of the report. Locked parts are removed server-side (not just hidden)."""
    locked = {
        "changes": max(0, len(a.changes) - 1),
        "missing_keywords": len(a.missing_keywords),
        "gaps": max(0, len(a.gaps) - 1),
        "strengths": max(0, len(a.strengths) - 1),
        "red_flags": len(a.red_flags),
        "requirements": len(a.requirements),
    }
    return a.model_copy(update={
        "changes": a.changes[:1],
        "missing_keywords": [],
        "gaps": a.gaps[:1],
        "strengths": a.strengths[:1],
        "red_flags": [],
        "requirements": [],
        "categories": [c.model_copy(update={"rationale": "", "detail": ""}) for c in a.categories],
    }), locked


def _client_id(request: Request) -> str:
    ip = request.client.host if request.client else "unknown"
    return hashlib.sha256(f"{_SALT}:{ip}".encode()).hexdigest()[:24]


async def enforce_limit(request: Request, action: str):
    limit, group, what = {
        "chat": (CHAT_LIMIT, ("chat",), "messages"),
        "preview": (PREVIEW_LIMIT, ("preview",), "free previews"),
        "jobs": (JOBS_LIMIT, ("jobs",), "job searches"),
    }.get(action, (DAILY_LIMIT, ("analyze", "kit"), "runs"))
    if limit <= 0 or is_admin(request):
        return
    used = await cache.count_and_record_usage(_client_id(request), action, group=group)
    if used > limit:
        extra = f" For unlimited access, email {CONTACT_EMAIL}." if action == "preview" else ""
        raise HTTPException(429, f"You've reached today's limit of {limit} {what}. Try again tomorrow.{extra}")


@app.get("/api/status")
def status():
    return {
        "demo": ai.demo_mode(),
        "model": ai.model_name(),
        "linkedin_sources": linkedin_sources() or ["LinkedIn public page"],
        "storage": "supabase" if cache.supabase_enabled() else "disk",
        "access_code_required": bool(ACCESS_CODE),
        "preview_available": bool(ACCESS_CODE),
        "contact_email": CONTACT_EMAIL,
        "job_countries": jobs.COUNTRIES,
        "admin_login": bool(ADMIN_USER and ADMIN_PASSWORD),
    }


class AdminLogin(BaseModel):
    username: str
    password: str


def require_admin(request: Request):
    if not is_admin(request):
        raise HTTPException(403, "Sign in as admin to see saved resumes.")


async def remember_run(request: Request, resume_name: str, resume_text: str, job: Job, analysis: Analysis):
    """Admin runs go into the history pool (with the resume text) so they can be previewed later."""
    if not is_admin(request):
        return
    fingerprint = "\n".join([analysis.mode, resume_name, resume_text, job.url, job.description[:2000]])
    await cache.save_history(hashlib.sha256(fingerprint.encode()).hexdigest()[:32], {
        "resume_name": resume_name, "mode": analysis.mode, "job_title": analysis.job_title, "company": analysis.company,
        "job_url": job.url, "overall_score": analysis.overall_score, "standing": analysis.standing,
        "resume_text": resume_text, "job": job.model_dump(mode="json"), "analysis": analysis.model_dump(mode="json"),
    })


@app.get("/api/history", dependencies=[Depends(require_admin)])
async def history():
    items = await cache.list_history()
    # score trend per job: every run against the same job, oldest first, up to and including this one
    runs: dict[str, list] = {}
    for it in sorted(items, key=lambda i: i["last_used"] or ""):
        job_key = it.get("job_url") or f"{it.get('mode')}|{it.get('job_title')}|{it.get('company')}"
        runs.setdefault(job_key, []).append(it.get("overall_score"))
        it["trend"] = list(runs[job_key])
    return {"items": items}


class Outcome(BaseModel):
    outcome: str


@app.post("/api/history/{key}/outcome", dependencies=[Depends(require_admin)])
async def history_outcome(key: str, body: Outcome):
    if not re.fullmatch(r"[0-9a-f]{32}", key) or body.outcome not in cache.OUTCOMES:
        raise HTTPException(400, "Unknown run or outcome.")
    if not await cache.set_history_outcome(key, body.outcome):
        raise HTTPException(404, "That saved run is gone.")
    return {"ok": True}


@app.get("/api/history/{key}", dependencies=[Depends(require_admin)])
async def history_item(key: str):
    if not re.fullmatch(r"[0-9a-f]{32}", key):
        raise HTTPException(404, "Not found.")
    item = await cache.get_history(key)
    if not item:
        raise HTTPException(404, "That saved run is gone.")
    return {k: item.get(k) for k in ("resume_name", "resume_text", "job", "analysis")}


@app.get("/api/admin")
def admin_status(request: Request):
    return {"admin": is_admin(request)}


@app.post("/api/admin/login")
def admin_login(body: AdminLogin, request: Request, response: Response):
    if not (ADMIN_USER and ADMIN_PASSWORD):
        raise HTTPException(404, "Admin sign-in isn't set up on this server.")
    client = _client_id(request)
    if _login_failures.get(client, 0) >= 10:
        raise HTTPException(429, "Too many failed sign-ins. Try again later.")
    ok = hmac.compare_digest(body.username.strip().lower(), ADMIN_USER.lower()) & hmac.compare_digest(body.password, ADMIN_PASSWORD)
    if not ok:
        _login_failures[client] = _login_failures.get(client, 0) + 1
        raise HTTPException(401, "Wrong username or password.")
    _login_failures.pop(client, None)
    response.set_cookie(ADMIN_COOKIE, _admin_token(), max_age=ADMIN_MAX_AGE, httponly=True,
                        secure=request.url.scheme == "https" or request.headers.get("x-forwarded-proto") == "https",
                        samesite="strict")
    return {"admin": True}


@app.post("/api/admin/logout")
def admin_logout(response: Response):
    response.delete_cookie(ADMIN_COOKIE)
    return {"admin": False}


@app.post("/api/analyze")
async def analyze(
    request: Request,
    resume: UploadFile = File(...),
    job_url: str = Form(""),
    job_text: str = Form(""),
):
    data = await resume.read()
    if len(data) > MAX_UPLOAD:
        raise HTTPException(413, "Resume file is over 5 MB.")
    try:
        resume_text, file_warnings = extract_resume(resume.filename or "", data)
    except ValueError as e:
        raise HTTPException(400, str(e))

    full = has_full_access(request)
    if not full and request.headers.get("x-access-code", "").strip():
        raise HTTPException(401, {"code": "access_code", "message": "That access code isn't valid. Clear it for a free preview."})
    await enforce_limit(request, "analyze" if full else "preview")

    # No job given: score the resume on its own. With a job: the full match process.
    if job_text.strip():
        job = Job(url=job_url.strip(), description=job_text.strip())
    elif job_url.strip():
        try:
            job = await fetch_job(job_url, premium=full)
        except JobFetchError as e:
            raise HTTPException(422, {"code": "job_fetch_failed", "message": str(e)})
    else:
        job = Job(description="")

    try:
        analysis = await (ai.analyze(resume_text, job) if job.description else ai.analyze_resume(resume_text))
    except ai.AIError as e:
        raise HTTPException(502, str(e))
    analysis.warnings = file_warnings
    await cache.save_analysis({
        "job_url": job.url,
        "job_title": analysis.job_title,
        "company": analysis.company,
        "candidate_name": analysis.candidate_name,
        "overall_score": analysis.overall_score,
        "standing": analysis.standing,
        "must_pct": analysis.coverage.must_pct if analysis.coverage else None,
        "resume_sha256": hashlib.sha256(resume_text.encode()).hexdigest(),
        "demo": ai.demo_mode(),
        "job": job.model_dump(mode="json"),
        "analysis": analysis.model_dump(mode="json"),
    })
    await remember_run(request, resume.filename or "Resume", resume_text, job, analysis)
    locked = None
    if not full:
        analysis, locked = preview_of(analysis)
    return {"resume_text": resume_text, "job": job, "analysis": analysis, "demo": ai.demo_mode(), "preview": not full, "locked": locked}


@app.post("/api/linkedin")
async def linkedin_review(
    request: Request,
    profile_url: str = Form(""),
    target_role: str = Form(""),
    profile_pdf: UploadFile | None = File(None),
):
    """Evaluate a LinkedIn profile against a target title: from the profile link (scraped) or the
    'Save to PDF' export (free)."""
    full = has_full_access(request)
    if not full and request.headers.get("x-access-code", "").strip():
        raise HTTPException(401, {"code": "access_code", "message": "That access code isn't valid. Clear it for a free preview."})
    if not profile_url.strip() and not profile_pdf:
        raise HTTPException(400, "Add your LinkedIn profile link or upload your LinkedIn PDF.")
    await enforce_limit(request, "analyze" if full else "preview")
    try:
        if profile_pdf and profile_pdf.filename:
            data = await profile_pdf.read()
            if len(data) > MAX_UPLOAD:
                raise HTTPException(413, "That PDF is over 5 MB.")
            try:
                text = extract_resume_text(profile_pdf.filename, data)
            except ValueError as e:
                raise HTTPException(400, str(e))
            profile = await ai.profile_from_pdf_text(text)
        else:
            profile = await fetch_linkedin_profile(profile_url)
        analysis = await ai.analyze_profile(profile, target_role)
    except JobFetchError as e:
        raise HTTPException(422, {"code": "profile_fetch_failed", "message": str(e)})
    except ai.AIError as e:
        raise HTTPException(502, str(e))
    await cache.save_analysis({
        "job_url": profile_url.strip(),
        "job_title": analysis.job_title,
        "company": "",
        "candidate_name": analysis.candidate_name,
        "overall_score": analysis.overall_score,
        "standing": analysis.standing,
        "must_pct": None,
        "resume_sha256": hashlib.sha256(profile.model_dump_json().encode()).hexdigest(),
        "demo": ai.demo_mode(),
        "job": {"source": profile.source},
        "analysis": analysis.model_dump(mode="json"),
    })
    source = profile_pdf.filename if profile_pdf and profile_pdf.filename else profile_url.strip()
    await remember_run(request, f"LinkedIn: {profile.name or source}", "", Job(url=profile_url.strip(), description=""), analysis)
    locked = None
    if not full:
        analysis, locked = preview_of(analysis)
    return {"profile": {"source": profile.source, "url": profile_url.strip()}, "analysis": analysis,
            "demo": ai.demo_mode(), "preview": not full, "locked": locked}


class KitRequest(BaseModel):
    resume_text: str
    job: Job
    analysis: Analysis


@app.post("/api/compare", dependencies=[Depends(require_access)])
async def compare(
    request: Request,
    resumes: list[UploadFile] = File(...),
    job_url: str = Form(""),
    job_text: str = Form(""),
):
    """Premium: rank 2-4 resumes against the same job."""
    if not 2 <= len(resumes) <= 4:
        raise HTTPException(400, "Add 2 to 4 resumes to compare.")
    if not job_text.strip() and not job_url.strip():
        raise HTTPException(400, "Add the job link or paste the job description.")
    texts, names = [], []
    for f in resumes:
        data = await f.read()
        if len(data) > MAX_UPLOAD:
            raise HTTPException(413, f"{f.filename} is over 5 MB.")
        try:
            text, _ = extract_resume(f.filename or "", data)
        except ValueError as e:
            raise HTTPException(400, f"{f.filename}: {e}")
        name = f.filename or f"Resume {len(names) + 1}"
        while name in names:  # same file name twice: keep the columns distinguishable
            name += " (2)"
        texts.append(text); names.append(name)
    for _ in resumes:
        await enforce_limit(request, "analyze")
    if job_text.strip():
        job = Job(url=job_url.strip(), description=job_text.strip())
    else:
        try:
            job = await fetch_job(job_url, premium=True)
        except JobFetchError as e:
            raise HTTPException(422, {"code": "job_fetch_failed", "message": str(e)})
    try:
        result = await ai.compare_resumes(list(zip(names, texts)), job)
    except ai.AIError as e:
        raise HTTPException(502, str(e))
    return {"job": job, "comparison": result, "demo": ai.demo_mode()}


@app.post("/api/apply-kit", dependencies=[Depends(require_access)])
async def apply_kit(req: KitRequest, request: Request):
    if not req.job.description:
        raise HTTPException(400, "Add a job link first — the application kit is tailored to a specific job.")
    await enforce_limit(request, "kit")
    try:
        kit = await ai.build_kit(req.resume_text, req.job, req.analysis)
    except ai.AIError as e:
        raise HTTPException(502, str(e))
    return {"kit": kit, "demo": ai.demo_mode()}


class JobSearchRequest(BaseModel):
    role: str
    location: str = ""
    country: str = "USA"
    skills: list[str] = []  # skills the resume already shows; postings are sorted by how many of them they ask for


@app.post("/api/jobs", dependencies=[Depends(require_access)])
async def find_jobs(req: JobSearchRequest, request: Request):
    """Recent postings for the resume's best-fit role; the client runs the normal job match on the one picked."""
    if not req.role.strip():
        raise HTTPException(400, "Score your resume first so we know which role to search for.")
    if req.country not in jobs.COUNTRIES:
        raise HTTPException(400, "Pick a country from the list.")
    await enforce_limit(request, "jobs")
    try:
        found = await jobs.search(req.role, req.location, req.country)
    except jobs.JobSearchError as e:
        raise HTTPException(502, str(e))
    return {"jobs": jobs.rank_by_fit(found, req.skills)}


class ChatRequest(BaseModel):
    resume_text: str
    resume_markdown: str = ""
    messages: list[ChatMessage]
    job: Job | None = None
    analysis: Analysis | None = None


@app.post("/api/resume-chat", dependencies=[Depends(require_access)])
async def resume_chat(req: ChatRequest, request: Request):
    if not req.messages or req.messages[-1].role != "user" or not req.messages[-1].content.strip():
        raise HTTPException(400, "Type what you'd like changed.")
    if len(req.messages[-1].content) > 2000:
        raise HTTPException(400, "Keep requests under 2,000 characters.")
    await enforce_limit(request, "chat")
    try:
        edit = await ai.edit_resume(req.resume_text, req.resume_markdown, req.messages, req.job, req.analysis)
    except ai.AIError as e:
        raise HTTPException(502, str(e))
    return {"edit": edit, "demo": ai.demo_mode()}


class ExportRequest(BaseModel):
    markdown: str
    filename: str = "resume"


def _add_runs(par, text: str):
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if chunk:
            par.add_run(chunk).bold = i % 2 == 1


def _safe_name(filename: str) -> str:
    return re.sub(r"[^\w-]+", "_", filename).strip("_") or "resume"


@app.post("/api/export/docx", dependencies=[Depends(require_access)])
def export_docx(req: ExportRequest):
    doc = Document()
    for line in req.markdown.splitlines():
        s = line.rstrip()
        if not s.strip():
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", s)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)) - 1 if len(m.group(1)) > 1 else 0)
        elif re.match(r"^\s*[-*•]\s+", s):
            _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*•]\s+", "", s))
        else:
            _add_runs(doc.add_paragraph(), s)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{_safe_name(req.filename)}.docx"'},
    )


@app.post("/api/export/pdf", dependencies=[Depends(require_access)])
def export_pdf(req: ExportRequest):
    return Response(
        pdf.markdown_to_pdf(req.markdown),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{_safe_name(req.filename)}.pdf"'},
    )


def _asset_version(name: str) -> str:
    return hashlib.sha256((STATIC / name).read_bytes()).hexdigest()[:10]


@app.get("/")
def index():
    # Version the CSS/JS links by content so a deploy never leaves visitors on stale cached files.
    html = (STATIC / "index.html").read_text()
    for name in ("app.js", "style.css"):
        html = html.replace(f"/static/{name}", f"/static/{name}?v={_asset_version(name)}")
    return HTMLResponse(html, headers={"Cache-Control": "no-cache"})


app.mount("/static", StaticFiles(directory=STATIC), name="static")
